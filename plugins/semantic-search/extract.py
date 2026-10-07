import json
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path


def extract(filename):
    ext = Path(filename).suffix.lower()
    if ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(filename)
        return {'pages': [{'page': i + 1, 'text': p.extract_text() or ''} for i, p in enumerate(reader.pages)]}
    with zipfile.ZipFile(filename) as archive:
        if ext == '.docx':
            roots = ['word/document.xml']
        elif ext == '.pptx':
            roots = sorted((n for n in archive.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')), key=lambda n: int(Path(n).stem[5:]))
        elif ext == '.xlsx':
            from openpyxl import load_workbook
            book = load_workbook(filename, read_only=True, data_only=True)
            try:
                return {'pages': [{'page': i + 1, 'title': sheet.title, 'text': '\n'.join('\t'.join(str(v) if v is not None else '' for v in row) for row in sheet.iter_rows(values_only=True))} for i, sheet in enumerate(book.worksheets)]}
            finally:
                book.close()
        else:
            raise ValueError('不支持此文档格式')
        pages = []
        for i, name in enumerate(roots):
            root = ET.fromstring(archive.read(name))
            paragraphs = []
            for node in root.iter():
                if node.tag.rsplit('}', 1)[-1] == 'p':
                    text = ''.join(child.text or '' for child in node.iter() if child.tag.rsplit('}', 1)[-1] == 't')
                    if text:
                        paragraphs.append(text)
            pages.append({'page': i + 1 if ext == '.pptx' else None, 'text': '\n'.join(paragraphs)})
        return {'pages': pages}


try:
    print(json.dumps(extract(sys.argv[1]), ensure_ascii=False))
except Exception as exc:
    print(json.dumps({'error': str(exc)}, ensure_ascii=False))
    sys.exit(1)
